"""Durable row spooling and one-shot XLSX export for large collection tasks."""

from __future__ import annotations

import json
import os
import queue
import re
import sqlite3
import threading
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

from openpyxl import Workbook

from src.core.app_state import get_app_state_root
from src.core.xlsx import sanitize_xlsx_cell

SPOOL_SCHEMA_VERSION = 1
DEFAULT_SPOOL_QUEUE_SIZE = 16
DEFAULT_SPOOL_MAX_BATCH_ROWS = 1000
EXCEL_MAX_DATA_ROWS_PER_SHEET = 1_048_575


@dataclass(frozen=True)
class SpoolSheet:
    """Schema of one logical output sheet."""

    name: str
    fields: Sequence[str]

    def normalized_fields(self) -> tuple[str, ...]:
        return tuple(str(field) for field in self.fields)


@dataclass(frozen=True)
class SpoolRow:
    """One idempotently addressable output row."""

    row_key: str
    payload: Mapping[str, Any]
    item_order: int = 0
    row_order: int = 0


@dataclass
class _WriteItem:
    stage: str
    item_key: str
    rows_by_sheet: dict[str, list[SpoolRow]]
    meta: Mapping[str, Any]


@dataclass
class _WriteCommand:
    items: list[_WriteItem]
    done: threading.Event = field(default_factory=threading.Event)
    error: BaseException | None = None


@dataclass
class _FlushCommand:
    done: threading.Event = field(default_factory=threading.Event)
    error: BaseException | None = None


_STOP = object()
_SAFE_SEGMENT_RE = re.compile(r"[^A-Za-z0-9._-]+")


def _json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)


def _safe_segment(value: str) -> str:
    cleaned = _SAFE_SEGMENT_RE.sub("_", str(value or "")).strip("._")
    return cleaned[:80] or "task"


def build_spool_path(tool_id: str, fingerprint: str, run_suffix: str | None = None) -> Path:
    """Return the stable app-state path for one resumable collection spool."""

    filename = _safe_segment(fingerprint)
    if run_suffix:
        filename = f"{filename}_{_safe_segment(run_suffix)}"
    path = get_app_state_root() / "spools" / _safe_segment(tool_id) / f"{filename}.sqlite3"
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def find_latest_spool_path(tool_id: str, fingerprint: str) -> Path:
    """Return the newest resumable spool for a fingerprint, or its base path."""

    base = build_spool_path(tool_id, fingerprint)
    candidates = [candidate for candidate in (base, *base.parent.glob(f"{base.stem}_*.sqlite3")) if candidate.exists()]
    if not candidates:
        return base
    return max(candidates, key=lambda candidate: candidate.stat().st_mtime_ns)


def remove_spool_files(path: str | Path) -> bool:
    """Remove a spool and its SQLite sidecar files."""

    base = Path(path)
    candidates = (base, Path(f"{base}-wal"), Path(f"{base}-shm"))
    for attempt in range(3):
        for candidate in candidates:
            try:
                candidate.unlink(missing_ok=True)
            except OSError:
                pass
        if not any(candidate.exists() for candidate in candidates):
            return True
        if attempt < 2:
            time.sleep(0.05 * (attempt + 1))
    return False


class SqliteRowSpool:
    """A bounded, single-writer SQLite sink for rows produced by many threads."""

    def __init__(
        self,
        path: str | Path,
        sheets: Iterable[SpoolSheet],
        *,
        metadata: Mapping[str, Any] | None = None,
        queue_size: int = DEFAULT_SPOOL_QUEUE_SIZE,
        max_batch_rows: int = DEFAULT_SPOOL_MAX_BATCH_ROWS,
    ) -> None:
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.sheets = tuple(sheets)
        if not self.sheets:
            raise ValueError("At least one spool sheet is required")
        self._sheet_fields = {sheet.name: sheet.normalized_fields() for sheet in self.sheets}
        if len(self._sheet_fields) != len(self.sheets):
            raise ValueError("Spool sheet names must be unique")
        self.max_batch_rows = max(1, int(max_batch_rows))
        self._queue: queue.Queue[object] = queue.Queue(maxsize=max(1, int(queue_size)))
        self._closed = False
        self._fatal_error: BaseException | None = None
        self._initialize_database(metadata or {})
        self._writer_thread = threading.Thread(
            target=self._writer_loop,
            name=f"sqlite-row-spool-{self.path.stem}",
            daemon=False,
        )
        self._writer_thread.start()

    @staticmethod
    def _connect(path: Path) -> sqlite3.Connection:
        connection = sqlite3.connect(str(path), timeout=30.0)
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("PRAGMA synchronous=NORMAL")
        connection.execute("PRAGMA busy_timeout=30000")
        return connection

    def _initialize_database(self, metadata: Mapping[str, Any]) -> None:
        connection = self._connect(self.path)
        try:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS spool_meta (
                    key TEXT PRIMARY KEY,
                    value_json TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS spool_sheets (
                    position INTEGER NOT NULL,
                    sheet_name TEXT PRIMARY KEY,
                    fields_json TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS spool_rows (
                    sequence INTEGER PRIMARY KEY AUTOINCREMENT,
                    sheet_name TEXT NOT NULL,
                    row_key TEXT NOT NULL,
                    item_order INTEGER NOT NULL DEFAULT 0,
                    row_order INTEGER NOT NULL DEFAULT 0,
                    payload_json TEXT NOT NULL,
                    UNIQUE(sheet_name, row_key)
                );
                CREATE INDEX IF NOT EXISTS idx_spool_rows_export
                    ON spool_rows(sheet_name, item_order, row_order, sequence);
                CREATE TABLE IF NOT EXISTS spool_completed (
                    stage TEXT NOT NULL,
                    item_key TEXT NOT NULL,
                    meta_json TEXT NOT NULL,
                    completed_at REAL NOT NULL,
                    PRIMARY KEY(stage, item_key)
                );
                """
            )
            existing_version = connection.execute(
                "SELECT value_json FROM spool_meta WHERE key='schema_version'"
            ).fetchone()
            if existing_version is not None and int(json.loads(existing_version[0])) != SPOOL_SCHEMA_VERSION:
                raise ValueError(f"Unsupported spool schema: {existing_version[0]}")
            connection.execute(
                "INSERT OR REPLACE INTO spool_meta(key, value_json) VALUES (?, ?)",
                ("schema_version", _json_dumps(SPOOL_SCHEMA_VERSION)),
            )
            for key, value in metadata.items():
                connection.execute(
                    "INSERT OR REPLACE INTO spool_meta(key, value_json) VALUES (?, ?)",
                    (str(key), _json_dumps(value)),
                )
            existing_sheets = {
                row[0]: tuple(json.loads(row[1]))
                for row in connection.execute("SELECT sheet_name, fields_json FROM spool_sheets")
            }
            for position, sheet in enumerate(self.sheets):
                fields = sheet.normalized_fields()
                existing = existing_sheets.get(sheet.name)
                if existing is not None and existing != fields:
                    raise ValueError(f"Spool sheet schema mismatch: {sheet.name}")
                connection.execute(
                    "INSERT OR IGNORE INTO spool_sheets(position, sheet_name, fields_json) VALUES (?, ?, ?)",
                    (position, sheet.name, _json_dumps(fields)),
                )
            unknown_sheets = set(existing_sheets) - set(self._sheet_fields)
            if unknown_sheets:
                raise ValueError(f"Spool contains unexpected sheets: {sorted(unknown_sheets)}")
            connection.commit()
        finally:
            connection.close()

    def _writer_loop(self) -> None:
        connection: sqlite3.Connection | None = None
        try:
            connection = self._connect(self.path)
            while True:
                command = self._queue.get()
                try:
                    if command is _STOP:
                        return
                    if self._fatal_error is not None:
                        command.error = self._fatal_error  # type: ignore[attr-defined]
                        continue
                    if isinstance(command, _FlushCommand):
                        connection.commit()
                        continue
                    if not isinstance(command, _WriteCommand):
                        raise TypeError(f"Unknown spool command: {type(command)!r}")
                    self._execute_write(connection, command)
                except BaseException as exc:
                    try:
                        connection.rollback()
                    except sqlite3.Error:
                        pass
                    self._fatal_error = exc
                    if hasattr(command, "error"):
                        command.error = exc  # type: ignore[attr-defined]
                finally:
                    if hasattr(command, "done"):
                        command.done.set()  # type: ignore[attr-defined]
                    self._queue.task_done()
        except BaseException as exc:
            self._fatal_error = exc
        finally:
            if connection is not None:
                connection.close()

    @staticmethod
    def _execute_write(connection: sqlite3.Connection, command: _WriteCommand) -> None:
        connection.execute("BEGIN")
        for item in command.items:
            for sheet_name, rows in item.rows_by_sheet.items():
                for row in rows:
                    connection.execute(
                        """
                        INSERT INTO spool_rows(sheet_name, row_key, item_order, row_order, payload_json)
                        VALUES (?, ?, ?, ?, ?)
                        ON CONFLICT(sheet_name, row_key) DO UPDATE SET
                            item_order=excluded.item_order,
                            row_order=excluded.row_order,
                            payload_json=excluded.payload_json
                        """,
                        (
                            sheet_name,
                            row.row_key,
                            int(row.item_order),
                            int(row.row_order),
                            _json_dumps(dict(row.payload)),
                        ),
                    )
            connection.execute(
                """
                INSERT INTO spool_completed(stage, item_key, meta_json, completed_at)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(stage, item_key) DO UPDATE SET
                    meta_json=excluded.meta_json,
                    completed_at=excluded.completed_at
                """,
                (item.stage, item.item_key, _json_dumps(dict(item.meta)), time.time()),
            )
        connection.commit()

    def _raise_if_unavailable(self) -> None:
        if self._closed:
            raise RuntimeError("Spool is closed")
        if self._fatal_error is not None:
            raise RuntimeError("Spool writer failed") from self._fatal_error

    def _enqueue_and_wait(self, command: _WriteCommand | _FlushCommand) -> None:
        while True:
            self._raise_if_unavailable()
            if not self._writer_thread.is_alive():
                raise RuntimeError("Spool writer stopped unexpectedly") from self._fatal_error
            try:
                self._queue.put(command, timeout=0.2)
                break
            except queue.Full:
                continue
        while not command.done.wait(timeout=0.2):
            if not self._writer_thread.is_alive():
                raise RuntimeError("Spool writer stopped unexpectedly") from self._fatal_error
        if command.error is not None:
            raise RuntimeError("Failed to write spool command") from command.error
        self._raise_if_unavailable()

    def submit_item(
        self,
        stage: str,
        item_key: str,
        rows: Mapping[str, Iterable[SpoolRow]],
        meta: Mapping[str, Any] | None = None,
    ) -> None:
        """Commit all rows and completion metadata for one logical item atomically."""

        self.submit_items(((stage, item_key, rows, meta),))

    def submit_items(
        self,
        items: Iterable[
            tuple[str, str, Mapping[str, Iterable[SpoolRow]], Mapping[str, Any] | None]
        ],
    ) -> None:
        """Commit multiple logical items together, with at most 1000 rows per batch."""

        self._raise_if_unavailable()
        write_items: list[_WriteItem] = []
        total_rows = 0
        for stage, item_key, rows, meta in items:
            normalized_stage = str(stage or "").strip()
            normalized_item_key = str(item_key or "").strip()
            if not normalized_stage or not normalized_item_key:
                raise ValueError("stage and item_key are required")
            materialized: dict[str, list[SpoolRow]] = {}
            for sheet_name, sheet_rows in rows.items():
                if sheet_name not in self._sheet_fields:
                    raise KeyError(f"Unknown spool sheet: {sheet_name}")
                current = list(sheet_rows)
                for row in current:
                    if not str(row.row_key or "").strip():
                        raise ValueError(f"Empty row key in sheet: {sheet_name}")
                materialized[sheet_name] = current
                total_rows += len(current)
            write_items.append(_WriteItem(normalized_stage, normalized_item_key, materialized, dict(meta or {})))
        if not write_items:
            return
        if total_rows > self.max_batch_rows:
            raise ValueError(f"Spool item exceeds max batch rows: {total_rows} > {self.max_batch_rows}")
        command = _WriteCommand(write_items)
        self._enqueue_and_wait(command)

    def flush(self) -> None:
        self._raise_if_unavailable()
        command = _FlushCommand()
        self._enqueue_and_wait(command)

    def close(self) -> None:
        if self._closed:
            return
        error: BaseException | None = None
        try:
            self.flush()
        except BaseException as exc:
            error = exc
        self._closed = True
        if self._writer_thread.is_alive():
            self._queue.put(_STOP)
            self._writer_thread.join()
        if error is not None:
            raise error
        if self._fatal_error is not None:
            raise RuntimeError("Spool writer failed") from self._fatal_error

    def discard(self) -> None:
        try:
            self.close()
        finally:
            remove_spool_files(self.path)

    def is_item_completed(self, stage: str, item_key: str) -> bool:
        self.flush()
        connection = sqlite3.connect(str(self.path), timeout=30.0)
        try:
            row = connection.execute(
                "SELECT 1 FROM spool_completed WHERE stage=? AND item_key=?",
                (str(stage), str(item_key)),
            ).fetchone()
            return row is not None
        finally:
            connection.close()

    def completed_keys(self, stage: str) -> set[str]:
        self.flush()
        connection = sqlite3.connect(str(self.path), timeout=30.0)
        try:
            return {
                str(row[0])
                for row in connection.execute("SELECT item_key FROM spool_completed WHERE stage=?", (str(stage),))
            }
        finally:
            connection.close()

    def completed_items(self, stage: str) -> dict[str, dict[str, Any]]:
        """Return completion metadata keyed by item key for one stage."""

        self.flush()
        connection = sqlite3.connect(str(self.path), timeout=30.0)
        try:
            return {
                str(item_key): json.loads(meta_json)
                for item_key, meta_json in connection.execute(
                    "SELECT item_key, meta_json FROM spool_completed WHERE stage=?",
                    (str(stage),),
                )
            }
        finally:
            connection.close()

    def row_count(self, sheet_name: str | None = None) -> int:
        self.flush()
        connection = sqlite3.connect(str(self.path), timeout=30.0)
        try:
            if sheet_name is None:
                row = connection.execute("SELECT COUNT(*) FROM spool_rows").fetchone()
            else:
                row = connection.execute("SELECT COUNT(*) FROM spool_rows WHERE sheet_name=?", (sheet_name,)).fetchone()
            return int(row[0] if row else 0)
        finally:
            connection.close()

    def iter_rows(self, sheet_name: str) -> Iterable[tuple[str, dict[str, Any], int, int]]:
        """Yield stored rows in deterministic export order without retaining them in memory."""

        if sheet_name not in self._sheet_fields:
            raise KeyError(f"Unknown spool sheet: {sheet_name}")
        self.flush()
        connection = sqlite3.connect(str(self.path), timeout=30.0)
        try:
            cursor = connection.execute(
                """
                SELECT row_key, payload_json, item_order, row_order
                FROM spool_rows
                WHERE sheet_name=?
                ORDER BY item_order, row_order, sequence
                """,
                (sheet_name,),
            )
            for row_key, payload_json, item_order, row_order in cursor:
                yield str(row_key), json.loads(payload_json), int(item_order), int(row_order)
        finally:
            connection.close()

    def metadata(self, key: str, default: Any = None) -> Any:
        self.flush()
        connection = sqlite3.connect(str(self.path), timeout=30.0)
        try:
            row = connection.execute("SELECT value_json FROM spool_meta WHERE key=?", (str(key),)).fetchone()
            return json.loads(row[0]) if row is not None else default
        finally:
            connection.close()


def _safe_sheet_title(seed: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/:*?\[\]]", "", str(seed or "")).strip() or "数据"
    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate in used:
        tail = f"_{suffix}"
        candidate = f"{base[: 31 - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def export_spool_to_xlsx(
    spool_path: str | Path,
    output_path: str | Path,
    progress_callback: Callable[[str], None] | None = None,
    *,
    max_data_rows_per_sheet: int = EXCEL_MAX_DATA_ROWS_PER_SHEET,
) -> dict[str, Any]:
    """Stream a spool into a new XLSX workbook and serialize it exactly once."""

    source = Path(spool_path)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    row_limit = max(1, min(int(max_data_rows_per_sheet), EXCEL_MAX_DATA_ROWS_PER_SHEET))
    connection = sqlite3.connect(str(source), timeout=30.0)
    workbook = Workbook(write_only=True)
    used_titles: set[str] = set()
    exported_rows = 0
    sheet_parts = 0
    stage_counts: dict[str, int] = {}
    sheet_row_counts: dict[str, int] = {}
    started = time.monotonic()
    last_progress = started
    try:
        sheet_rows = list(
            connection.execute(
                "SELECT sheet_name, fields_json FROM spool_sheets ORDER BY position, sheet_name"
            )
        )
        if not sheet_rows:
            raise ValueError(f"Spool has no sheet schema: {source}")
        stage_counts = {
            str(stage): int(count)
            for stage, count in connection.execute(
                "SELECT stage, COUNT(*) FROM spool_completed GROUP BY stage ORDER BY stage"
            )
        }
        sheet_row_counts = {
            str(sheet_name): int(count)
            for sheet_name, count in connection.execute(
                "SELECT sheet_name, COUNT(*) FROM spool_rows GROUP BY sheet_name ORDER BY sheet_name"
            )
        }
        for logical_name, fields_json in sheet_rows:
            fields = tuple(str(value) for value in json.loads(fields_json))
            part_number = 1
            current_count = 0
            title_seed = str(logical_name)
            worksheet = workbook.create_sheet(_safe_sheet_title(title_seed, used_titles))
            worksheet.append([sanitize_xlsx_cell(field) for field in fields])
            sheet_parts += 1
            cursor = connection.execute(
                """
                SELECT payload_json
                FROM spool_rows
                WHERE sheet_name=?
                ORDER BY item_order, row_order, sequence
                """,
                (logical_name,),
            )
            for (payload_json,) in cursor:
                if current_count >= row_limit:
                    part_number += 1
                    worksheet = workbook.create_sheet(
                        _safe_sheet_title(f"{logical_name}_{part_number}", used_titles)
                    )
                    worksheet.append([sanitize_xlsx_cell(field) for field in fields])
                    current_count = 0
                    sheet_parts += 1
                payload = json.loads(payload_json)
                worksheet.append([sanitize_xlsx_cell(payload.get(field, "")) for field in fields])
                current_count += 1
                exported_rows += 1
                now = time.monotonic()
                if progress_callback and (exported_rows % 50_000 == 0 or now - last_progress >= 5.0):
                    progress_callback(f"XLSX 导出进度：已写入 {exported_rows} 行。")
                    last_progress = now
        temp_path = Path(f"{target}.tmp")
        try:
            workbook.save(temp_path)
            os.replace(temp_path, target)
        except Exception:
            if progress_callback and temp_path.exists():
                progress_callback(f"XLSX 原子替换失败，临时文件已保留：{temp_path}")
            raise
    finally:
        connection.close()
    elapsed = time.monotonic() - started
    if progress_callback:
        progress_callback(f"XLSX 导出完成：{exported_rows} 行，{sheet_parts} 个工作表，耗时 {elapsed:.1f} 秒。")
    return {
        "exported_rows": exported_rows,
        "sheet_parts": sheet_parts,
        "elapsed_seconds": elapsed,
        "output_path": str(target),
        "stage_counts": stage_counts,
        "sheet_row_counts": sheet_row_counts,
    }
