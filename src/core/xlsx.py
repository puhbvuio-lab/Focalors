"""
Excel (XLSX) 文件处理工具，使用 openpyxl 库提供流式行写入支持，具备防注入机制与原子写入保护。
"""

from __future__ import annotations

import os
import re
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from src.core.csv_utils import sanitize_csv_cell

DEFAULT_XLSX_AUTOSAVE_ROWS = 5000
_XLSX_ILLEGAL_CHAR_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")


def sanitize_xlsx_cell(value: Any) -> Any:
    """
    清洗并格式化单个 Excel 单元格的值。
    防范 CSV/Excel 注入漏洞：在 Excel 中，以 =, +, -, @ 开头的内容会被误识别为公式并执行。
    如果在其头部追加单引号 ' 则强制 Excel 将其作为纯文本显示。
    """
    value = sanitize_csv_cell(value)
    if isinstance(value, str):
        value = _XLSX_ILLEGAL_CHAR_RE.sub("", value)
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _apply_wrap_alignment(worksheet, fieldnames: Iterable[str], wrap_fields: Iterable[str]) -> None:
    """
    给刚写入的一行里指定字段开启「自动换行」，使单元格内的换行真正换行显示。

    仅在存在 wrap_fields 时生效，避免给所有表格增加无意义的样式开销。
    """
    wrap_set = set(wrap_fields or ())
    if not wrap_set:
        return
    row_index = worksheet.max_row
    for column_index, name in enumerate(fieldnames, start=1):
        if name in wrap_set:
            worksheet.cell(row=row_index, column=column_index).alignment = Alignment(wrap_text=True, vertical="top")


class XlsxRowWriter:
    """
    单表格 Excel 行写入器，适合流式追加单 sheet 数据。
    """

    def __init__(
        self,
        output_path: str,
        fieldnames: Iterable[str],
        sheet_name: str = "数据",
        autosave_every: int = DEFAULT_XLSX_AUTOSAVE_ROWS,
        append: bool = False,
        wrap_fields: Iterable[str] = (),
    ):
        """
        Args:
            output_path: 输出的 Excel 文件路径
            fieldnames: 表头字段列表
            sheet_name: 表格的工作表名称
            autosave_every: 累计写入多少行时执行一次磁盘保存（防数据丢失）
            append: 是否追加模式
            wrap_fields: 需要开启「自动换行」的字段名（单元格内含换行时逐行显示）
        """
        self.output_path = str(output_path)
        Path(self.output_path).parent.mkdir(parents=True, exist_ok=True)
        self.fieldnames = list(fieldnames)
        self.wrap_fields = set(wrap_fields or ())
        self.autosave_every = max(1, int(autosave_every or DEFAULT_XLSX_AUTOSAVE_ROWS))
        self._rows_since_save = 0
        self._save_count = 0
        if append and Path(self.output_path).exists():
            self._load_for_append(sheet_name)
        else:
            self.workbook = Workbook()
            self.worksheet = self.workbook.active
            # Excel 规定 Sheet 名称最大长度为 31 字符，超出直接被截断
            self.worksheet.title = sheet_name[:31] or "数据"
            self.worksheet.append(self.fieldnames)
            self.save()

    def _load_for_append(self, sheet_name: str) -> None:
        workbook = load_workbook(self.output_path)
        expected_sheet_name = sheet_name[:31] or "数据"
        worksheet = workbook[expected_sheet_name] if expected_sheet_name in workbook.sheetnames else workbook.active
        header = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1), [])]
        if header != self.fieldnames:
            raise ValueError(f"Existing XLSX header does not match: {self.output_path}")
        self.workbook = workbook
        self.worksheet = worksheet

    def writerow(self, row: Mapping[str, Any]):
        """
        写入单行数据。
        """
        self._append_row(row)
        self._rows_since_save += 1
        if self._rows_since_save >= self.autosave_every:
            self.save()

    def writerows(self, rows: Iterable[Mapping[str, Any]]):
        """
        批量写入多行数据。
        """
        for row in rows:
            self._append_row(row)
            self._rows_since_save += 1
            if self._rows_since_save >= self.autosave_every:
                self.save()

    def _append_row(self, row: Mapping[str, Any]):
        self.worksheet.append([sanitize_xlsx_cell(row.get(field, "")) for field in self.fieldnames])
        _apply_wrap_alignment(self.worksheet, self.fieldnames, self.wrap_fields)

    def save(self):
        """
        原子式保存当前 Excel 文件，写入临时文件再执行 os.replace，
        防止保存过程中程序被异常终止或断电导致原文件损坏。
        """
        temp_path = f"{self.output_path}.tmp"
        self.workbook.save(temp_path)
        try:
            os.replace(temp_path, self.output_path)
        except OSError:
            # 在某些 Windows 网络共享磁盘或者并发占用下，os.replace 可能会报 OSError，
            # 此时 fallback 回直接保存
            self.workbook.save(self.output_path)
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except OSError:
                    pass
        self._rows_since_save = 0
        self._save_count += 1

class MultiSheetXlsxWriter:
    """
    多表格 Excel 写入器，支持同时在单个工作簿下写入多个 Sheet。
    """

    def __init__(
        self,
        output_path: str,
        sheets_fields: dict[str, list[str]],
        autosave_every: int = DEFAULT_XLSX_AUTOSAVE_ROWS,
        append: bool = False,
        sheets_wrap_fields: dict[str, Iterable[str]] | None = None,
    ):
        """
        Args:
            output_path: 输出的 Excel 文件路径
            sheets_fields: 键值对，键为 sheet_name，值为表头字段列表
            autosave_every: 缓存多少行时自动保存
            append: 是否追加模式
            sheets_wrap_fields: 每个 sheet 需要开启「自动换行」的字段名
        """
        self.output_path = str(output_path)
        Path(self.output_path).parent.mkdir(parents=True, exist_ok=True)
        self.sheets_fields = sheets_fields
        self.sheets_wrap_fields = {name: set(fields or ()) for name, fields in (sheets_wrap_fields or {}).items()}
        self.autosave_every = max(1, int(autosave_every or DEFAULT_XLSX_AUTOSAVE_ROWS))
        self._rows_since_save = 0
        self._save_count = 0

        if append and Path(self.output_path).exists():
            self._load_for_append()
        else:
            self.workbook = Workbook()
            # 移除 openpyxl 实例化时自动创建的默认 Sheet，以便仅保留用户自定义的 Sheet
            default_sheet = self.workbook.active
            if default_sheet is not None:
                self.workbook.remove(default_sheet)

            self.worksheets = {}
            for sheet_name, fieldnames in sheets_fields.items():
                # 同样遵循 31 字符上限限制
                ws = self.workbook.create_sheet(title=sheet_name[:31] or "Sheet")
                ws.append(list(fieldnames))
                self.worksheets[sheet_name] = ws
            self.save()

    def _load_for_append(self) -> None:
        workbook = load_workbook(self.output_path)
        worksheets = {}
        changed = False
        requested_sheet_names = {sheet_name[:31] or "Sheet" for sheet_name in self.sheets_fields}
        if requested_sheet_names.isdisjoint(set(workbook.sheetnames)):
            has_existing_data = any(ws.max_row > 1 or ws.max_column > 1 for ws in workbook.worksheets)
            if has_existing_data:
                raise ValueError(f"Existing XLSX has no requested sheets and appears to use a different schema: {self.output_path}")
        for sheet_name, fieldnames in self.sheets_fields.items():
            excel_sheet_name = sheet_name[:31] or "Sheet"
            if excel_sheet_name in workbook.sheetnames:
                ws = workbook[excel_sheet_name]
                header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1), [])]
                if header != list(fieldnames):
                    raise ValueError(f"Existing XLSX header does not match sheet '{sheet_name}': {self.output_path}")
            else:
                ws = workbook.create_sheet(title=excel_sheet_name)
                ws.append(list(fieldnames))
                changed = True
            worksheets[sheet_name] = ws
        self.workbook = workbook
        self.worksheets = worksheets
        if changed:
            self.save()

    def writerow(self, sheet_name: str, row: Mapping[str, Any]):
        """
        向指定名称的工作表中追加单行数据。
        """
        if sheet_name not in self.worksheets:
            import logging
            logging.getLogger(__name__).warning("writerow: sheet '%s' not registered, row skipped", sheet_name)
            return
        fieldnames = self.sheets_fields[sheet_name]
        ws = self.worksheets[sheet_name]
        ws.append([sanitize_xlsx_cell(row.get(field, "")) for field in fieldnames])
        _apply_wrap_alignment(ws, fieldnames, self.sheets_wrap_fields.get(sheet_name, ()))
        self._rows_since_save += 1
        if self._rows_since_save >= self.autosave_every:
            self.save()

    def writerows(self, sheet_name: str, rows: Iterable[Mapping[str, Any]]):
        """
        向指定名称的工作表中批量追加多行数据。
        """
        if sheet_name not in self.worksheets:
            import logging
            logging.getLogger(__name__).warning("writerows: sheet '%s' not registered, rows skipped", sheet_name)
            return
        fieldnames = self.sheets_fields[sheet_name]
        ws = self.worksheets[sheet_name]
        for row in rows:
            ws.append([sanitize_xlsx_cell(row.get(field, "")) for field in fieldnames])
            _apply_wrap_alignment(ws, fieldnames, self.sheets_wrap_fields.get(sheet_name, ()))
            self._rows_since_save += 1
            if self._rows_since_save >= self.autosave_every:
                self.save()

    def save(self):
        """
        原子式保存当前工作薄文件。
        """
        temp_path = f"{self.output_path}.tmp"
        self.workbook.save(temp_path)
        try:
            os.replace(temp_path, self.output_path)
        except OSError:
            self.workbook.save(self.output_path)
            if os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except OSError:
                    pass
        self._rows_since_save = 0
        self._save_count += 1
