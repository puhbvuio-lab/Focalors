from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl


def normalize_header(value: Any) -> str:
    return str(value or "").strip()


def validate_headers(header_values: list[Any]) -> list[str]:
    headers = [normalize_header(value) for value in header_values]
    if not headers or not any(headers):
        raise ValueError("目标 sheet 的表头为空，无法选择判定列。")
    duplicates: set[str] = set()
    seen: set[str] = set()
    for header in headers:
        if not header:
            raise ValueError("目标 sheet 存在空表头，请先补齐表头。")
        lowered = header.casefold()
        if lowered in seen:
            duplicates.add(header)
        seen.add(lowered)
    if duplicates:
        duplicate_text = "、".join(sorted(duplicates))
        raise ValueError(f"目标 sheet 存在重复表头：{duplicate_text}")
    return headers


def resolve_target_columns(headers: list[str], selected_columns: list[str]) -> list[str]:
    selected = {normalize_header(value).casefold() for value in selected_columns if normalize_header(value)}
    ordered = [header for header in headers if header.casefold() in selected]
    if not ordered:
        raise ValueError("至少需要选择 1 个判定列。")
    if len(ordered) != len(selected):
        resolved = {header.casefold() for header in ordered}
        missing = [value for value in selected_columns if normalize_header(value).casefold() not in resolved]
        raise ValueError(f"存在无效的判定列：{', '.join(missing)}")
    return ordered


def extract_sheet_names(input_xlsx: str | Path) -> list[str]:
    workbook = openpyxl.load_workbook(input_xlsx, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def extract_sheet_headers(input_xlsx: str | Path, sheet_name: str) -> list[str]:
    workbook = openpyxl.load_workbook(input_xlsx, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet 不存在：{sheet_name}")
        worksheet = workbook[sheet_name]
        return validate_headers([cell.value for cell in worksheet[1]])
    finally:
        workbook.close()
