from __future__ import annotations

import os
import sqlite3

from openpyxl import load_workbook
import pytest

from src.core.row_spool import (
    SpoolRow,
    SpoolSheet,
    SqliteRowSpool,
    build_spool_path,
    export_spool_to_xlsx,
    find_latest_spool_path,
)


def test_spool_is_idempotent_and_exports_in_stable_order(tmp_path):
    spool_path = tmp_path / "rows.sqlite3"
    output_path = tmp_path / "rows.xlsx"
    spool = SqliteRowSpool(
        spool_path,
        [SpoolSheet("视频信息", ["编号", "内容"]), SpoolSheet("评论信息", ["编号", "内容"])],
    )
    spool.submit_item(
        "video",
        "b",
        {"视频信息": [SpoolRow("b", {"编号": "2", "内容": "second"}, 2, 0)]},
    )
    spool.submit_item(
        "video",
        "a",
        {"视频信息": [SpoolRow("a", {"编号": "1", "内容": "first"}, 1, 0)]},
    )
    spool.submit_item(
        "video",
        "a",
        {"视频信息": [SpoolRow("a", {"编号": "1", "内容": "updated"}, 1, 0)]},
    )
    spool.close()

    stats = export_spool_to_xlsx(spool_path, output_path)
    assert stats["exported_rows"] == 2
    assert stats["stage_counts"] == {"video": 2}
    assert stats["sheet_row_counts"] == {"视频信息": 2}
    workbook = load_workbook(output_path, read_only=True)
    assert list(workbook["视频信息"].iter_rows(values_only=True)) == [
        ("编号", "内容"),
        ("1", "updated"),
        ("2", "second"),
    ]
    assert list(workbook["评论信息"].iter_rows(values_only=True)) == [("编号", "内容")]
    workbook.close()


def test_spool_export_splits_large_logical_sheet(tmp_path):
    spool_path = tmp_path / "split.sqlite3"
    output_path = tmp_path / "split.xlsx"
    spool = SqliteRowSpool(spool_path, [SpoolSheet("评论信息", ["值"])])
    for index in range(5):
        spool.submit_item(
            "comment",
            str(index),
            {"评论信息": [SpoolRow(str(index), {"值": index}, index, 0)]},
        )
    spool.close()

    export_spool_to_xlsx(spool_path, output_path, max_data_rows_per_sheet=2)
    workbook = load_workbook(output_path, read_only=True)
    assert workbook.sheetnames == ["评论信息", "评论信息_2", "评论信息_3"]
    assert [len(list(ws.iter_rows(values_only=True))) for ws in workbook.worksheets] == [3, 3, 2]
    workbook.close()


def test_spool_export_sanitizes_formula_and_illegal_characters(tmp_path):
    spool_path = tmp_path / "sanitize.sqlite3"
    output_path = tmp_path / "sanitize.xlsx"
    spool = SqliteRowSpool(spool_path, [SpoolSheet("数据", ["值"])])
    spool.submit_item(
        "item",
        "1",
        {"数据": [SpoolRow("1", {"值": "=cmd\x02"})]},
    )
    spool.close()

    export_spool_to_xlsx(spool_path, output_path)
    workbook = load_workbook(output_path, read_only=True)
    assert workbook["数据"]["A2"].value == "'=cmd"
    workbook.close()


def test_spool_persists_completed_items_for_resume(tmp_path):
    spool_path = tmp_path / "resume.sqlite3"
    schema = [SpoolSheet("数据", ["值"])]
    first = SqliteRowSpool(spool_path, schema)
    first.submit_item("video", "abc", {"数据": [SpoolRow("abc", {"值": 1})]})
    first.close()

    resumed = SqliteRowSpool(spool_path, schema)
    assert resumed.is_item_completed("video", "abc") is True
    assert resumed.completed_keys("video") == {"abc"}
    assert resumed.completed_items("video") == {"abc": {}}
    assert resumed.row_count("数据") == 1
    resumed.close()


def test_find_latest_spool_prefers_newest_isolated_run(monkeypatch, tmp_path):
    monkeypatch.setattr("src.core.row_spool.get_app_state_root", lambda: tmp_path)
    base = build_spool_path("tool", "fingerprint")
    isolated = build_spool_path("tool", "fingerprint", "run-2")
    base.touch()
    isolated.touch()
    os.utime(base, (1, 1))
    os.utime(isolated, (2, 2))

    assert find_latest_spool_path("tool", "fingerprint") == isolated


def test_spool_enforces_bounded_batch_before_transaction(tmp_path):
    spool_path = tmp_path / "bounded.sqlite3"
    spool = SqliteRowSpool(
        spool_path,
        [SpoolSheet("数据", ["值"])],
        queue_size=16,
        max_batch_rows=2,
    )
    assert spool._queue.maxsize == 16
    with pytest.raises(ValueError, match="exceeds max batch rows"):
        spool.submit_items(
            (
                ("item", "first", {"数据": [SpoolRow("1", {"值": 1})]}, None),
                ("item", "second", {"数据": [SpoolRow(str(index), {"值": index}) for index in range(2, 4)]}, None),
            )
        )
    spool.close()

    with sqlite3.connect(spool_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM spool_rows").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM spool_completed").fetchone()[0] == 0


def test_spool_writer_failure_propagates_without_hanging(monkeypatch, tmp_path):
    def fail_write(_connection, _command):
        raise sqlite3.OperationalError("disk failure")

    monkeypatch.setattr(SqliteRowSpool, "_execute_write", staticmethod(fail_write))
    spool = SqliteRowSpool(tmp_path / "failure.sqlite3", [SpoolSheet("数据", ["值"])])
    with pytest.raises(RuntimeError, match="Failed to write spool command"):
        spool.submit_item("item", "1", {"数据": [SpoolRow("1", {"值": 1})]})
    with pytest.raises(RuntimeError, match="Spool writer failed"):
        spool.close()
