"""X 推文详情采集：媒体链接、正文链接还原与评论扫描停止条件的回归测试。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.core.csv_utils import MultilineText, sanitize_csv_cell
from src.core.xlsx import MultiSheetXlsxWriter, XlsxRowWriter, sanitize_xlsx_cell
from src.platforms.x_twitter.comments import (
    _count_first_level_articles,
    _has_recommendation_boundary,
    should_stop_scanning,
)
from src.platforms.x_twitter.tweet_media import (
    build_syndication_token,
    extract_media_from_syndication,
    extract_quoted_tweet_id,
    pick_best_video_variant,
    upgrade_twitter_image_url,
)
from src.platforms.x_twitter.tweet_metrics import (
    CSV_FIELDS,
    MEDIA_FIELD_IMAGE,
    MEDIA_FIELD_VIDEO,
    WRAP_FIELDS,
    normalize_tweet_text,
    wait_for_article_media_ready,
)


class FakePage:
    """只实现被测函数用到的最小接口。"""

    def __init__(self, evaluate_result):
        self._result = evaluate_result
        self.scripts: list[str] = []

    def evaluate(self, script):
        self.scripts.append(script)
        return self._result


class TestSyndicationToken(unittest.TestCase):
    def test_token_matches_known_values(self):
        # 这两个值来自真实接口调用（2026-09 实测可用）
        self.assertEqual(build_syndication_token("2096455855851933861"), "52y7kkgz4ztq8q6fn6lmcxr")
        self.assertEqual(build_syndication_token("2098762850311295223"), "535ghi7s7t3rym4zwzsemi")

    def test_token_has_no_zero_or_dot(self):
        token = build_syndication_token("1")
        self.assertNotIn("0", token)
        self.assertNotIn(".", token)


class TestMediaParsing(unittest.TestCase):
    def test_photos_are_deduped_and_upgraded(self):
        payload = {
            "mediaDetails": [
                {"type": "photo", "media_url_https": "https://pbs.twimg.com/media/AAA.jpg"},
                {"type": "photo", "media_url_https": "https://pbs.twimg.com/media/AAA.jpg"},
                {"type": "photo", "media_url_https": "https://pbs.twimg.com/media/BBB?format=jpg&name=small"},
            ]
        }
        images, videos = extract_media_from_syndication(payload)
        self.assertEqual(
            images,
            [
                "https://pbs.twimg.com/media/AAA.jpg",
                "https://pbs.twimg.com/media/BBB?format=jpg&name=large",
            ],
        )
        self.assertEqual(videos, [])

    def test_best_mp4_variant_is_selected(self):
        variants = [
            {"type": "application/x-mpegURL", "src": "https://video.twimg.com/a/pl/x.m3u8?tag=16"},
            {"type": "video/mp4", "src": "https://video.twimg.com/a/vid/avc1/480x270/low.mp4?tag=16"},
            {"type": "video/mp4", "src": "https://video.twimg.com/a/vid/avc1/1280x720/high.mp4?tag=16"},
        ]
        self.assertEqual(
            pick_best_video_variant(variants),
            "https://video.twimg.com/a/vid/avc1/1280x720/high.mp4?tag=16",
        )

    def test_video_media_detail_yields_video_and_poster(self):
        payload = {
            "mediaDetails": [
                {
                    "type": "video",
                    "media_url_https": "https://pbs.twimg.com/media/poster.jpg",
                    "video_info": {
                        "variants": [
                            {"type": "video/mp4", "src": "https://video.twimg.com/amplify_video/1/vid/avc1/640x360/a.mp4?tag=16"},
                            {"type": "video/mp4", "src": "https://video.twimg.com/amplify_video/1/vid/avc1/1920x1080/b.mp4?tag=16"},
                        ]
                    },
                }
            ]
        }
        images, videos = extract_media_from_syndication(payload)
        self.assertEqual(images, ["https://pbs.twimg.com/media/poster.jpg"])
        self.assertEqual(videos, ["https://video.twimg.com/amplify_video/1/vid/avc1/1920x1080/b.mp4?tag=16"])

    def test_top_level_video_is_used(self):
        payload = {"video": {"variants": [{"type": "video/mp4", "src": "https://video.twimg.com/x/vid/avc1/320x180/a.mp4"}]}}
        _, videos = extract_media_from_syndication(payload)
        self.assertEqual(videos, ["https://video.twimg.com/x/vid/avc1/320x180/a.mp4"])

    def test_missing_payload_returns_empty(self):
        self.assertEqual(extract_media_from_syndication(None), ([], []))
        self.assertEqual(extract_media_from_syndication({}), ([], []))

    def test_image_url_upgrade_strips_thumbnail_params(self):
        # 带扩展名：去掉查询串即原图
        self.assertEqual(
            upgrade_twitter_image_url("https://pbs.twimg.com/media/HSJPfz5X0AIx-Ou.jpg?name=small"),
            "https://pbs.twimg.com/media/HSJPfz5X0AIx-Ou.jpg",
        )
        self.assertEqual(
            upgrade_twitter_image_url("https://pbs.twimg.com/media/HSAfpmmXgAcbHXz.jpg"),
            "https://pbs.twimg.com/media/HSAfpmmXgAcbHXz.jpg",
        )

    def test_image_url_without_extension_keeps_format(self):
        # 无扩展名：查询串必须保留，否则实测 404
        self.assertEqual(
            upgrade_twitter_image_url("https://pbs.twimg.com/media/HSGaryAWQAI9Bhg?format=jpg&name=small"),
            "https://pbs.twimg.com/media/HSGaryAWQAI9Bhg?format=jpg&name=large",
        )

    def test_image_url_other_host_untouched(self):
        self.assertEqual(upgrade_twitter_image_url("https://example.com/a.jpg"), "https://example.com/a.jpg")
        self.assertEqual(upgrade_twitter_image_url(""), "")


class TestTweetTextNormalization(unittest.TestCase):
    def test_quoted_tweet_id_is_extracted(self):
        payload = {"quoted_tweet": {"id_str": "2097882762426618362"}}
        self.assertEqual(extract_quoted_tweet_id(payload), "2097882762426618362")

    def test_quoted_tweet_id_missing(self):
        self.assertEqual(extract_quoted_tweet_id(None), "")
        self.assertEqual(extract_quoted_tweet_id({}), "")
        self.assertEqual(extract_quoted_tweet_id({"quoted_tweet": {}}), "")

    def test_broken_url_is_rejoined(self):
        raw = "List of winners:\n\nhttps://\nhoyo.link/0vTJjwvZk\n※ rewards"
        self.assertIn("https://hoyo.link/0vTJjwvZk", normalize_tweet_text(raw))

    def test_space_broken_url_is_rejoined(self):
        self.assertEqual(normalize_tweet_text("see https:// hoyo.link/abc now"), "see https://hoyo.link/abc now")

    def test_plain_text_untouched(self):
        text = "Hello Travelers!\nSee you in Teyvat."
        self.assertEqual(normalize_tweet_text(text), text)

    def test_empty_text(self):
        self.assertEqual(normalize_tweet_text(""), "")


class TestMultilineCell(unittest.TestCase):
    def test_multiline_text_keeps_newlines(self):
        value = MultilineText("https://img/1.jpg\nhttps://img/2.jpg")
        self.assertEqual(sanitize_csv_cell(value), "https://img/1.jpg\nhttps://img/2.jpg")
        self.assertEqual(sanitize_xlsx_cell(value), "https://img/1.jpg\nhttps://img/2.jpg")

    def test_plain_text_still_flattens(self):
        self.assertEqual(sanitize_csv_cell("a\nb"), "a b")
        self.assertEqual(sanitize_xlsx_cell("a\nb"), "a b")


class TestTweetMetricsSchema(unittest.TestCase):
    def test_media_columns_are_declared(self):
        self.assertIn(MEDIA_FIELD_IMAGE, CSV_FIELDS)
        self.assertIn(MEDIA_FIELD_VIDEO, CSV_FIELDS)
        self.assertEqual(CSV_FIELDS[-2:], [MEDIA_FIELD_IMAGE, MEDIA_FIELD_VIDEO])
        self.assertEqual(set(WRAP_FIELDS), {MEDIA_FIELD_IMAGE, MEDIA_FIELD_VIDEO})

    def test_single_sheet_writer_applies_wrap_and_keeps_newlines(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "single.xlsx")
            writer = XlsxRowWriter(path, CSV_FIELDS, wrap_fields=WRAP_FIELDS)
            writer.writerow(
                {
                    "序号": "1",
                    "推文链接": "https://x.com/a/status/1",
                    MEDIA_FIELD_IMAGE: MultilineText("https://img/1.jpg\nhttps://img/2.jpg"),
                    MEDIA_FIELD_VIDEO: MultilineText("https://video/x.mp4"),
                }
            )
            writer.save()

            worksheet = load_workbook(path).active
            image_cell = worksheet.cell(row=2, column=CSV_FIELDS.index(MEDIA_FIELD_IMAGE) + 1)
            self.assertEqual(image_cell.value, "https://img/1.jpg\nhttps://img/2.jpg")
            self.assertTrue(image_cell.alignment.wrap_text)
            text_cell = worksheet.cell(row=2, column=CSV_FIELDS.index("推文链接") + 1)
            self.assertFalse(text_cell.alignment.wrap_text)

    def test_multi_sheet_writer_applies_wrap_on_tweet_sheet_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "multi.xlsx")
            comment_fields = ["序号", "推文链接", "评论的点赞量", "评论内容", "评论发布时间"]
            writer = MultiSheetXlsxWriter(
                path,
                {"推文信息": CSV_FIELDS, "评论信息": comment_fields},
                sheets_wrap_fields={"推文信息": WRAP_FIELDS},
            )
            writer.writerow("推文信息", {"序号": "1", MEDIA_FIELD_IMAGE: MultilineText("https://img/1.jpg\nhttps://img/2.jpg")})
            writer.writerow("评论信息", {"序号": "1", "评论内容": "a\nb"})
            writer.save()

            workbook = load_workbook(path)
            tweet_sheet = workbook["推文信息"]
            image_cell = tweet_sheet.cell(row=2, column=CSV_FIELDS.index(MEDIA_FIELD_IMAGE) + 1)
            self.assertEqual(image_cell.value, "https://img/1.jpg\nhttps://img/2.jpg")
            self.assertTrue(image_cell.alignment.wrap_text)

            comment_sheet = workbook["评论信息"]
            self.assertEqual(comment_sheet.cell(row=2, column=4).value, "a b")
            self.assertFalse(comment_sheet.cell(row=2, column=4).alignment.wrap_text)


class TestMediaReadyWait(unittest.TestCase):
    class _Article:
        def __init__(self, signatures):
            self._signatures = list(signatures)
            self.calls = 0

        def evaluate(self, script):
            index = min(self.calls, len(self._signatures) - 1)
            self.calls += 1
            return self._signatures[index]

    def test_returns_after_signature_stabilizes(self):
        article = self._Article(
            [
                {"images": 0, "videos": 1, "shells": 2},
                {"images": 1, "videos": 1, "shells": 2},
                {"images": 1, "videos": 1, "shells": 2},
                {"images": 1, "videos": 1, "shells": 2},
            ]
        )
        signature = wait_for_article_media_ready(article, timeout=5.0, poll=0.0)
        self.assertEqual(signature, {"images": 1, "videos": 1, "shells": 2})

    def test_returns_last_signature_when_never_stable(self):
        article = self._Article([{"images": 0, "videos": 0, "shells": 1}])
        signature = wait_for_article_media_ready(article, timeout=0.05, poll=0.01)
        self.assertEqual(signature.get("shells"), 1)

    def test_tolerates_detached_article(self):
        class Detached:
            def evaluate(self, script):
                raise RuntimeError("detached")

        self.assertEqual(wait_for_article_media_ready(Detached(), timeout=0.05, poll=0.0), {})


class TestCommentScanStopRule(unittest.TestCase):
    def test_boundary_alone_does_not_stop(self):
        # 回归：历史实现一看到「发现更多」就结束，导致抓到 0～8 条评论
        stop, _ = should_stop_scanning(0, 6, True, 1, 3)
        self.assertFalse(stop)

    def test_stops_after_no_new_limit(self):
        stop, reason = should_stop_scanning(6, 6, False, 0, 3)
        self.assertTrue(stop)
        self.assertIn("没有发现新评论", reason)

    def test_stops_after_boundary_patience(self):
        stop, reason = should_stop_scanning(3, 6, True, 3, 3)
        self.assertTrue(stop)
        self.assertIn("推荐区", reason)

    def test_continues_while_new_comments_arrive(self):
        stop, _ = should_stop_scanning(0, 6, True, 0, 3)
        self.assertFalse(stop)

    def test_count_first_level_articles(self):
        page = FakePage(7)
        self.assertEqual(_count_first_level_articles(page), 7)
        self.assertIn('tabindex="0"', page.scripts[0])

    def test_count_first_level_articles_tolerates_failure(self):
        class BrokenPage:
            def evaluate(self, script):
                raise RuntimeError("detached")

        self.assertEqual(_count_first_level_articles(BrokenPage()), 0)

    def test_boundary_detection_returns_true_when_marker_present(self):
        self.assertTrue(_has_recommendation_boundary(FakePage(True)))

    def test_boundary_detection_tolerates_failure(self):
        class BrokenPage:
            def evaluate(self, script):
                raise RuntimeError("detached")

        self.assertFalse(_has_recommendation_boundary(BrokenPage()))


if __name__ == "__main__":
    unittest.main(verbosity=2)
