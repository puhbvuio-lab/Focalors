from openpyxl import load_workbook

from src.core.xlsx import DEFAULT_XLSX_AUTOSAVE_ROWS, MultiSheetXlsxWriter, XlsxRowWriter


def test_default_xlsx_autosave_threshold_is_5000(tmp_path):
    writer = XlsxRowWriter(str(tmp_path / "default.xlsx"), ["value"])
    assert writer.autosave_every == DEFAULT_XLSX_AUTOSAVE_ROWS == 5000
    assert writer._save_count == 1


def test_row_writer_only_saves_at_threshold_or_explicit_flush(tmp_path):
    output_path = tmp_path / "rows.xlsx"
    writer = XlsxRowWriter(str(output_path), ["value"], autosave_every=3)
    initial_saves = writer._save_count

    writer.writerows([{"value": 1}, {"value": 2}])
    assert writer._save_count == initial_saves
    writer.writerow({"value": 3})
    assert writer._save_count == initial_saves + 1
    writer.writerow({"value": 4})
    writer.save()
    assert writer._save_count == initial_saves + 2

    rows = list(load_workbook(output_path).active.iter_rows(values_only=True))
    assert rows == [("value",), (1,), (2,), (3,), (4,)]


def test_multi_sheet_writer_counts_rows_across_sheets_without_per_batch_save(tmp_path):
    output_path = tmp_path / "multi.xlsx"
    writer = MultiSheetXlsxWriter(
        str(output_path),
        {"first": ["value"], "second": ["value"]},
        autosave_every=4,
    )
    initial_saves = writer._save_count

    writer.writerows("first", [{"value": 1}, {"value": 2}])
    writer.writerow("second", {"value": 3})
    assert writer._save_count == initial_saves
    writer.writerow("second", {"value": 4})
    assert writer._save_count == initial_saves + 1

    workbook = load_workbook(output_path)
    assert list(workbook["first"].iter_rows(values_only=True)) == [("value",), (1,), (2,)]
    assert list(workbook["second"].iter_rows(values_only=True)) == [("value",), (3,), (4,)]
    workbook.close()


def test_xlsx_writer_strips_excel_illegal_control_characters(tmp_path):
    output_path = tmp_path / "illegal_chars.xlsx"
    writer = XlsxRowWriter(str(output_path), ["text"])

    writer.writerow({"text": "mini\x02map"})
    writer.save()

    workbook = load_workbook(output_path)
    assert workbook.active["A2"].value == "minimap"
    workbook.close()
