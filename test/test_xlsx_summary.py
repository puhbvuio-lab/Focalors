from openpyxl import Workbook, load_workbook

from src.core import summarize_outputs


def _write_workbook(path, title: str, value: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    worksheet.append(["value"])
    worksheet.append([value])
    workbook.save(path)


def test_summarize_outputs_is_exported_and_combines_workbooks(tmp_path):
    first_path = tmp_path / "first.xlsx"
    second_path = tmp_path / "second.xlsx"
    _write_workbook(first_path, "first", "alpha")
    _write_workbook(second_path, "second", "beta")

    summary_path = summarize_outputs([first_path, second_path])

    assert summary_path is not None
    workbook = load_workbook(summary_path, read_only=True, data_only=True)
    try:
        assert len(workbook.sheetnames) == 2
        values = {
            cell
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows(min_row=2, values_only=True)
            for cell in row
            if cell
        }
        assert values == {"alpha", "beta"}
    finally:
        workbook.close()
