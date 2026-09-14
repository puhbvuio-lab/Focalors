from __future__ import annotations

import sqlite3
import threading
from contextlib import closing

from openpyxl import load_workbook

from src.core import RunStatus, TaskFeedback
from src.platforms.youtube import channel_works as channel_module
from src.platforms.youtube.comments import CommentFetchResult


class DummyPool:
    def __init__(self, api_keys):
        self.api_keys = api_keys


def _work(video_id: str, comments: str = "1") -> dict[str, str]:
    return {
        "link": f"https://www.youtube.com/watch?v={video_id}",
        "title": f"Video {video_id}",
        "channel_title": "Channel",
        "channel_id": "UC1",
        "published_at": "2026-07-03 10:00:00",
        "video_type": "视频",
        "views": "10",
        "likes": "2",
        "comments": comments,
    }


def _configure(monkeypatch, tmp_path):
    output_path = tmp_path / "youtube_channel_works.xlsx"
    spool_path = tmp_path / "youtube_channel_works.sqlite3"
    monkeypatch.setattr(channel_module, "build_output_path", lambda *_args, **_kwargs: str(output_path))
    monkeypatch.setattr(channel_module, "channel_works_spool_path", lambda _scope: str(spool_path))
    monkeypatch.setattr(channel_module, "YouTubeClientPool", DummyPool)
    return output_path, spool_path


def _run(finish_callback, stop_event=None, feedback=None):
    config = {"verify_video_type": "否", "max_parallel_tabs": 1}
    if feedback is not None:
        config["_task_feedback"] = feedback
    return channel_module.run_youtube_channel_works_spider(
        ["key"],
        "https://www.youtube.com/@channel",
        collect_target="仅视频与Shorts",
        max_video_items=100,
        max_post_scrolls=10,
        fetch_shorts_related="否",
        live_stream_policy="不处理",
        limit_time_str="否",
        get_comments_str="是",
        max_comments=10,
        log_callback=lambda _message: None,
        finish_callback=finish_callback,
        stop_event=stop_event,
        config=config,
    )


def test_channel_works_writes_works_before_bounded_comment_collection(monkeypatch, tmp_path):
    output_path, spool_path = _configure(monkeypatch, tmp_path)
    monkeypatch.setattr(
        channel_module,
        "collect_video_works_with_api",
        lambda *_args, **_kwargs: [_work("workvideo01", "0"), _work("workvideo02", "2")],
    )
    fetched = []

    def fake_fetch(_api_keys, tasks, *_args, result_callback=None, retain_results=True, **_kwargs):
        assert retain_results is False
        with closing(sqlite3.connect(spool_path)) as connection:
            assert connection.execute("SELECT COUNT(*) FROM spool_rows WHERE sheet_name='作品信息'").fetchone()[0] == 2
        for task in tasks:
            fetched.append(task.video_id)
            result_callback(
                task,
                CommentFetchResult(
                    task.video_id,
                    [{"comment_id": "comment-1", "like_count": 3, "text": "ok", "published_at": ""}],
                    "ok",
                ),
            )
        return {}

    monkeypatch.setattr(channel_module, "fetch_top_comments_for_videos", fake_fetch)
    finished = []
    feedback = TaskFeedback()
    outcome = _run(finished.append, feedback=feedback)

    assert outcome.status == RunStatus.SUCCEEDED
    assert finished == [outcome]
    assert fetched == ["workvideo02"]
    assert output_path.exists()
    assert not spool_path.exists()
    assert feedback.snapshot()["completed"] == feedback.snapshot()["total"] == 2
    assert feedback.snapshot()["activity"] == "导出完成"
    workbook = load_workbook(output_path, read_only=True)
    assert len(list(workbook["作品信息"].iter_rows(values_only=True))) == 3
    assert len(list(workbook["评论信息"].iter_rows(values_only=True))) == 2
    workbook.close()


def test_channel_works_cancel_keeps_spool_and_does_not_export(monkeypatch, tmp_path):
    output_path, spool_path = _configure(monkeypatch, tmp_path)
    stop_event = threading.Event()

    def fake_collect(*_args, **_kwargs):
        stop_event.set()
        return [_work("workvideo01")]

    monkeypatch.setattr(channel_module, "collect_video_works_with_api", fake_collect)
    outcome = _run(lambda _outcome: None, stop_event)

    assert outcome.status == RunStatus.CANCELLED
    assert spool_path.exists()
    assert not output_path.exists()
    assert not (tmp_path / "youtube_channel_works.xlsx.tmp").exists()


def test_channel_works_numbering_follows_input_order_under_concurrency(monkeypatch, tmp_path):
    output_path, _spool_path = _configure(monkeypatch, tmp_path)
    second_started = threading.Event()

    def fake_collect(_pool, channel_url, *_args, **_kwargs):
        if channel_url.endswith("@first"):
            assert second_started.wait(timeout=2)
            return [_work("firstvideo1")]
        second_started.set()
        return [_work("secondvideo")]

    monkeypatch.setattr(channel_module, "collect_video_works_with_api", fake_collect)
    outcome = channel_module.run_youtube_channel_works_spider(
        ["key"],
        "https://www.youtube.com/@first\nhttps://www.youtube.com/@second",
        collect_target="仅视频与Shorts",
        get_comments_str="否",
        log_callback=lambda _message: None,
        finish_callback=lambda _outcome: None,
        config={"verify_video_type": "否", "max_parallel_tabs": 2},
    )

    assert outcome.status == RunStatus.SUCCEEDED
    workbook = load_workbook(output_path, read_only=True)
    rows = list(workbook["作品信息"].iter_rows(values_only=True))
    header = {name: index for index, name in enumerate(rows[0])}
    assert [row[header["编号"]] for row in rows[1:]] == ["1", "2"]
    assert [row[header["作者主页链接"]] for row in rows[1:]] == [
        "https://www.youtube.com/@first",
        "https://www.youtube.com/@second",
    ]
    workbook.close()
