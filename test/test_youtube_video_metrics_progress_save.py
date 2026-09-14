# -*- coding: utf-8 -*-

from __future__ import annotations

import sqlite3
import threading

from openpyxl import load_workbook

from src.core import RunStatus, SpoolRow, SpoolSheet, SqliteRowSpool, TaskFeedback
from src.platforms.youtube import comments as comments_module
from src.platforms.youtube.comments import CommentFetchResult


class DummyPool:
    def __init__(self, api_keys):
        self.api_keys = api_keys


def _metric(video_id: str, comment_count: str = "1") -> dict:
    return {
        "标题": f"Video {video_id}",
        "频道名称": "Channel",
        "频道ID": "UC1",
        "发布日期": "2026-07-03 10:00:00",
        "直播状态": "",
        "视频时长": "00:01:00",
        "视频简介": "desc",
        "播放量": "10",
        "点赞数": "2",
        "评论数": comment_count,
    }


def _configure_paths(monkeypatch, tmp_path):
    output_path = tmp_path / "youtube_video_metrics.xlsx"
    spool_path = tmp_path / "youtube_video_metrics.sqlite3"
    monkeypatch.setattr(comments_module, "build_output_path", lambda *_args, **_kwargs: str(output_path))
    monkeypatch.setattr(comments_module, "video_metrics_spool_path", lambda _scope: str(spool_path))
    monkeypatch.setattr(comments_module, "YouTubeClientPool", DummyPool)
    return output_path, spool_path


def test_video_metrics_persists_video_stage_before_comment_fetch(monkeypatch, tmp_path):
    txt_path = tmp_path / "videos.txt"
    txt_path.write_text(
        "\n".join(
            [
                "https://www.youtube.com/watch?v=video000001",
                "https://www.youtube.com/watch?v=video000002",
            ]
        ),
        encoding="utf-8",
    )
    output_path, spool_path = _configure_paths(monkeypatch, tmp_path)
    monkeypatch.setattr(
        comments_module,
        "fetch_video_metrics",
        lambda _pool, video_ids, *_args, **_kwargs: {video_id: _metric(video_id) for video_id in video_ids},
    )

    def fake_fetch_comments(_api_keys, comment_tasks, *_args, result_callback=None, **_kwargs):
        connection = sqlite3.connect(spool_path)
        try:
            assert connection.execute("SELECT COUNT(*) FROM spool_rows WHERE sheet_name='视频信息'").fetchone()[0] == 2
        finally:
            connection.close()
        for task in comment_tasks:
            result_callback(
                task,
                CommentFetchResult(
                    task.video_id,
                    [{"comment_id": f"comment-{task.video_id}", "like_count": 1, "text": "comment", "published_at": ""}],
                    "ok",
                ),
            )
        return {}

    monkeypatch.setattr(comments_module, "fetch_top_comments_for_videos", fake_fetch_comments)
    finished = []
    feedback = TaskFeedback()
    outcome = comments_module.run_youtube_video_metrics_spider(
        ["key"],
        str(txt_path),
        "否",
        "不处理",
        "是",
        "否",
        100,
        lambda _message: None,
        finished.append,
        config={"comment_top_limit": 10, "max_parallel_tabs": 1, "_task_feedback": feedback},
    )

    assert outcome.status == RunStatus.SUCCEEDED
    assert finished == [outcome]
    assert output_path.exists()
    assert not spool_path.exists()
    assert feedback.snapshot()["completed"] == feedback.snapshot()["total"] == 5
    assert feedback.snapshot()["activity"] == "导出完成"
    workbook = load_workbook(output_path, read_only=True)
    assert len(list(workbook["视频信息"].iter_rows(values_only=True))) == 3
    assert len(list(workbook["评论信息"].iter_rows(values_only=True))) == 3
    workbook.close()


def test_video_metrics_type_detection_uses_fifty_video_batches(monkeypatch, tmp_path):
    video_ids = [f"video{i:06d}" for i in range(1, 52)]
    txt_path = tmp_path / "videos.txt"
    txt_path.write_text("\n".join(f"https://www.youtube.com/watch?v={video_id}" for video_id in video_ids), encoding="utf-8")
    _configure_paths(monkeypatch, tmp_path)
    monkeypatch.setattr(
        comments_module,
        "fetch_video_metrics",
        lambda _pool, ids, *_args, **_kwargs: {video_id: _metric(video_id) for video_id in ids},
    )
    type_batches = []
    persist_batches = []
    original_submit_items = comments_module.SqliteRowSpool.submit_items

    def spy_submit_items(self, items):
        materialized = list(items)
        persist_batches.append(len(materialized))
        return original_submit_items(self, materialized)

    def fake_check_type(ids, *_args, **_kwargs):
        type_batches.append(list(ids))
        return {video_id: "视频" for video_id in ids}

    monkeypatch.setattr(comments_module, "check_video_type_bulk", fake_check_type)
    monkeypatch.setattr(comments_module.SqliteRowSpool, "submit_items", spy_submit_items)
    outcome = comments_module.run_youtube_video_metrics_spider(
        ["key"], str(txt_path), "否", "不处理", "否", "是", 100,
        lambda _message: None, lambda _outcome: None, config={"max_parallel_tabs": 1},
    )

    assert outcome.status == RunStatus.SUCCEEDED
    assert [len(batch) for batch in type_batches] == [50, 1]
    assert persist_batches == [50, 1]


def test_video_metrics_skips_comment_fetch_when_comment_count_is_zero(monkeypatch, tmp_path):
    txt_path = tmp_path / "videos.txt"
    txt_path.write_text(
        "\n".join(
            [
                "https://www.youtube.com/watch?v=video000001",
                "https://www.youtube.com/watch?v=video000002",
            ]
        ),
        encoding="utf-8",
    )
    output_path, _spool_path = _configure_paths(monkeypatch, tmp_path)

    def fake_metrics(_pool, video_ids, *_args, **_kwargs):
        return {
            video_id: _metric(video_id, "0" if video_id == "video000001" else "5")
            for video_id in video_ids
        }

    fetched = []

    def fake_comments(_api_keys, tasks, *_args, result_callback=None, **_kwargs):
        for task in tasks:
            fetched.append(task.video_id)
            result_callback(task, CommentFetchResult(task.video_id, [], "empty"))
        return {}

    monkeypatch.setattr(comments_module, "fetch_video_metrics", fake_metrics)
    monkeypatch.setattr(comments_module, "fetch_top_comments_for_videos", fake_comments)
    outcome = comments_module.run_youtube_video_metrics_spider(
        ["key"], str(txt_path), "否", "不处理", "是", "否", 100,
        lambda _message: None, lambda _outcome: None, config={"max_parallel_tabs": 1},
    )

    assert outcome.status == RunStatus.SUCCEEDED
    assert fetched == ["video000002"]
    workbook = load_workbook(output_path, read_only=True)
    comment_rows = list(workbook["评论信息"].iter_rows(values_only=True))
    assert len(comment_rows) == 3
    workbook.close()


def test_video_metrics_cancel_keeps_spool_without_export(monkeypatch, tmp_path):
    txt_path = tmp_path / "videos.txt"
    txt_path.write_text("https://www.youtube.com/watch?v=video000001", encoding="utf-8")
    output_path, spool_path = _configure_paths(monkeypatch, tmp_path)
    stop_event = threading.Event()

    def fake_metrics(_pool, video_ids, *_args, **_kwargs):
        stop_event.set()
        return {video_id: _metric(video_id) for video_id in video_ids}

    monkeypatch.setattr(comments_module, "fetch_video_metrics", fake_metrics)
    outcome = comments_module.run_youtube_video_metrics_spider(
        ["key"], str(txt_path), "否", "不处理", "否", "否", 100,
        lambda _message: None, lambda _outcome: None, stop_event=stop_event, config={"max_parallel_tabs": 1},
    )

    assert outcome.status == RunStatus.CANCELLED
    assert spool_path.exists()
    assert not output_path.exists()
    assert not (tmp_path / "youtube_video_metrics.xlsx.tmp").exists()


def test_video_metrics_resumes_completed_video_stage_without_duplicates(monkeypatch, tmp_path):
    txt_path = tmp_path / "videos.txt"
    txt_path.write_text(
        "https://www.youtube.com/watch?v=video000001\nhttps://www.youtube.com/watch?v=video000002",
        encoding="utf-8",
    )
    output_path, spool_path = _configure_paths(monkeypatch, tmp_path)
    first_row = {field: "" for field in comments_module.VIDEO_FIELDS}
    first_row.update({"编号": "1", "视频链接": "https://www.youtube.com/watch?v=video000001", "标题": "cached"})
    spool = SqliteRowSpool(spool_path, [SpoolSheet("视频信息", comments_module.VIDEO_FIELDS)])
    spool.submit_item(
        "video",
        "video000001",
        {"视频信息": [SpoolRow("video000001", first_row, 1, 0)]},
    )
    spool.close()
    requested_batches = []

    def fake_metrics(_pool, video_ids, *_args, **_kwargs):
        requested_batches.append(list(video_ids))
        return {video_id: _metric(video_id) for video_id in video_ids}

    monkeypatch.setattr(comments_module, "fetch_video_metrics", fake_metrics)
    outcome = comments_module.run_youtube_video_metrics_spider(
        ["key"], str(txt_path), "否", "不处理", "否", "否", 100,
        lambda _message: None, lambda _outcome: None, config={"max_parallel_tabs": 1},
    )

    assert outcome.status == RunStatus.SUCCEEDED
    assert requested_batches == [["video000002"]]
    workbook = load_workbook(output_path, read_only=True)
    rows = list(workbook["视频信息"].iter_rows(values_only=True))
    assert len(rows) == 3
    assert [row[0] for row in rows[1:]] == ["1", "2"]
    workbook.close()


def test_video_metrics_explicit_partial_export_keeps_resume_spool(monkeypatch, tmp_path):
    txt_path = tmp_path / "videos.txt"
    txt_path.write_text("https://www.youtube.com/watch?v=video000001", encoding="utf-8")
    output_path, spool_path = _configure_paths(monkeypatch, tmp_path)
    row = {field: "" for field in comments_module.VIDEO_FIELDS}
    row.update({"编号": "1", "视频链接": "https://www.youtube.com/watch?v=video000001"})
    spool = SqliteRowSpool(spool_path, [SpoolSheet("视频信息", comments_module.VIDEO_FIELDS)])
    spool.submit_item("video", "video000001", {"视频信息": [SpoolRow("video000001", row, 1, 0)]})
    spool.close()
    monkeypatch.setattr(
        comments_module,
        "fetch_video_metrics",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("partial export must not collect")),
    )

    outcome = comments_module.run_youtube_video_metrics_spider(
        ["key"], str(txt_path), "否", "不处理", "否", "否", 100,
        lambda _message: None, lambda _outcome: None, config={"_resume_action": "export"},
    )

    assert outcome.status == RunStatus.PARTIAL
    assert outcome.output_path == str(output_path)
    assert output_path.exists()
    assert spool_path.exists()


def test_video_metrics_recoverable_network_abort_returns_partial_and_keeps_spool(monkeypatch, tmp_path):
    txt_path = tmp_path / "videos.txt"
    txt_path.write_text("https://www.youtube.com/watch?v=video000001", encoding="utf-8")
    output_path, spool_path = _configure_paths(monkeypatch, tmp_path)
    monkeypatch.setattr(
        comments_module,
        "fetch_video_metrics",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ConnectionResetError(10054, "connection reset")),
    )

    outcome = comments_module.run_youtube_video_metrics_spider(
        ["key"], str(txt_path), "否", "不处理", "否", "否", 100,
        lambda _message: None, lambda _outcome: None, config={"max_parallel_tabs": 1},
    )

    assert outcome.status == RunStatus.PARTIAL
    assert outcome.output_path is None
    assert outcome.errors[0].code == "YOUTUBE_RECOVERABLE_STOPPED"
    assert spool_path.exists()
    assert not output_path.exists()
